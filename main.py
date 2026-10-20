#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Keyword-based extraction from SPSS-exported Excel workbooks.

Public demonstration refactored from a Colab research prototype.
It does NOT reproduce or validate the original study's statistics.
"""
from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path
from typing import Any

KEYWORDS = ("百分比", "次數", "顯著性。", "Exp(B)", "上界", "下界")
EXPECTED_LENGTHS = (2, 6)


def load_sheets(path: Path) -> dict[str, list[list[Any]]]:
    """Read workbooks as two-dimensional lists, keeping cell positions."""
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return {
                ws.title: [list(row) for row in ws.iter_rows(values_only=True)]
                for ws in wb.worksheets
            }
        finally:
            wb.close()
    if path.suffix.lower() == '.xls':
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError('Reading .xls requires xlrd: pip install xlrd') from exc
        book = xlrd.open_workbook(str(path))
        return {
            ws.name: [ws.row_values(i) for i in range(ws.nrows)]
            for ws in book.sheets()
        }
    raise ValueError(f'Unsupported file type: {path.suffix}')


def cell(rows: list[list[Any]], row: int, col: int) -> Any:
    if row < 0 or row >= len(rows) or col < 0 or col >= len(rows[row]):
        return None
    return rows[row][col]


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def extract_records(
    rows: list[list[Any]],
    filename: str,
    worksheet: str,
    max_rows: int = 8,
) -> list[dict[str, Any]]:
    """Find matching headers, collect nonempty cells beneath them.

    Original prototype relied on the shape of cells above/below each
    keyword. This demonstration retains the idea, reports shapes, and
    flags mismatches for human review rather than guessing their meaning.
    """
    occurrences: Counter[str] = Counter()
    records: list[dict[str, Any]] = []
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if not isinstance(value, str) or value.strip() not in KEYWORDS:
                continue
            keyword = value.strip()
            occurrences[keyword] += 1
            raw_window = [cell(rows, r+offset, c) for offset in range(1, max_rows+1)]
            extracted = [v for v in raw_window if not is_blank(v)]
            above = cell(rows, r-1, c)
            shape = len(extracted)
            status = 'candidate' if shape in EXPECTED_LENGTHS else 'review_needed'
            records.append({
                'file': filename,
                'sheet': worksheet,
                'keyword': keyword,
                'occurrence': occurrences[keyword],
                'excel_row': r+1,
                'excel_col': c+1,
                'cell_above': '' if is_blank(above) else str(above),
                'count': shape,
                'shape': f'{shape}-value',
                'status': status,
                'values': extracted,
            })
    return records


def latex_escape(text: Any) -> str:
    """Escape LaTeX-special characters in one pass."""
    subst = {'\\': r'\textbackslash{}', '&': r'\&', '%': r'\%', '$': r'\$',
             '#': r'\#', '_': r'\_', '{': r'\{', '}': r'\}',
             '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}
    return ''.join(subst.get(c, c) for c in str(text))


def export_results(records: list[dict[str, Any]], destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    columns = ('file', 'sheet', 'keyword', 'occurrence', 'excel_row',
               'excel_col', 'cell_above', 'count', 'shape', 'status', 'values')
    with (destination/'extracted.csv').open('w', encoding='utf-8-sig', newline='') as out:
        writer = csv.DictWriter(out, fieldnames=columns)
        writer.writeheader()
        for rec in records:
            line = rec.copy()
            line['values'] = json.dumps(rec['values'], ensure_ascii=False)
            writer.writerow(line)
    with (destination/'extracted.json').open('w', encoding='utf-8') as out:
        json.dump(records, out, ensure_ascii=False, indent=2, default=str)
    # Small generic LaTeX example; not a recreation of medical-report template.
    lines = [
        r'\documentclass{article}',
        r'\usepackage[utf8]{inputenc}',
        r'\usepackage{longtable}',
        r'\begin{document}',
        r'\begin{longtable}{p{3cm}p{2cm}p{1cm}p{7cm}}',
        r'File & Keyword & Count & Extracted values \\',
        r'\hline',
    ]
    for rec in records:
        fields = [latex_escape(rec['file']), latex_escape(rec['keyword']),
                  str(rec['count']), latex_escape(', '.join(map(str, rec['values'])))]
        lines.append(' & '.join(fields) + r' \\')
    lines.extend([r'\end{longtable}', r'\end{document}'])
    (destination/'report.tex').write_text('\n'.join(lines)+'\n', encoding='utf-8')
    warning_lines = [
        f"{r['file']} / {r['sheet']} / {r['keyword']} #{r['occurrence']}: "
        f"{r['count']} values (expected 2 or 6)"
        for r in records if r['status'] == 'review_needed'
    ]
    (destination/'needs_review.txt').write_text(
        '\n'.join(warning_lines) + ('\n' if warning_lines else 'No shape warnings. Manual verification is still required.\n'),
        encoding='utf-8',
    )


def run(input_dir: Path, output_dir: Path, max_rows: int = 8) -> list[dict[str, Any]]:
    files = sorted(p for p in input_dir.iterdir()
                   if p.is_file() and p.suffix.lower() in ('.xls', '.xlsx'))
    if not files:
        raise ValueError(f'No .xls/.xlsx workbooks found in {input_dir}')
    results: list[dict[str, Any]] = []
    for path in files:
        for sheet_name, rows in load_sheets(path).items():
            results.extend(extract_records(rows, path.name, sheet_name, max_rows))
    export_results(results, output_dir)
    return results


def main() -> None:
    parser = argparse.ArgumentParser(description='SPSS-exported Excel keyword extractor (demo)')
    parser.add_argument('--input', type=Path, default=Path('examples'), help='Input folder')
    parser.add_argument('--output', type=Path, default=Path('output'), help='Output folder')
    parser.add_argument('--max-rows', type=int, default=8, help='Cells to inspect below each header')
    args = parser.parse_args()
    if not args.input.is_dir():
        parser.error(f'Input directory does not exist: {args.input}')
    if args.max_rows < 1:
        parser.error('--max-rows must be a positive integer')
    results = run(args.input, args.output, args.max_rows)
    flagged = sum(rec['status'] == 'review_needed' for rec in results)
    print(f'Extracted {len(results)} keyword blocks; {flagged} require review.')
    print(f'Output: {args.output.resolve()}')
    print('Verify extracted values against the source workbook before using them in reports.')


if __name__ == '__main__':
    main()
