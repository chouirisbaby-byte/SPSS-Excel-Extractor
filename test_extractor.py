import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook

from main import extract_records, load_sheets, run, latex_escape


class TestExtractor(unittest.TestCase):
    def test_keyword_and_occurrence(self):
        rows = [['名稱', '次數'], ['範例', 10], ['', 20], ['', 30], ['', 40], ['', 50], ['', 60]]
        data = extract_records(rows, 'synthetic.xlsx', 'test', max_rows=6)
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['count'], 6)
        self.assertEqual(data[0]['values'], [10, 20, 30, 40, 50, 60])
        self.assertEqual(data[0]['status'], 'candidate')

    def test_length_mismatch_is_flagged(self):
        rows = [['上界'], [1.2], [2.3], [3.4]]
        rec = extract_records(rows, 'synthetic.xlsx', 'test', max_rows=3)[0]
        self.assertEqual(rec['status'], 'review_needed')

    def test_run_generates_outputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t/'input').mkdir()
            wb = Workbook()
            ws = wb.active
            ws.title = 'Demo'
            ws['B2'] = 'Exp(B)'
            ws['B3'] = 1.1
            ws['B4'] = 2.2
            wb.save(t/'input'/'fake.xlsx')
            records = run(t/'input', t/'out', max_rows=2)
            self.assertEqual(records[0]['values'], [1.1, 2.2])
            for name in ('extracted.csv', 'extracted.json', 'report.tex', 'needs_review.txt'):
                self.assertTrue((t/'out'/name).exists())

    def test_latex_escape(self):
        self.assertEqual(latex_escape('a_b&c'), r'a\_b\&c')


if __name__ == '__main__':
    unittest.main()
